"""Build the irrigation water consumption tables of Pfister & Bayer (2014) (decision L44).

Source: Pfister S., Bayer P. (2014) Monthly water stress: spatially and temporally explicit consumptive water
footprint of global crop production. Journal of Cleaner Production 73, 52-62, doi:10.1016/j.jclepro.2013.11.031
(file Literature/Pfister_JCP_inpress.pdf). Blue water consumption of irrigation in m3 per t of harvested crop,
input data of the year 2000; files supplied by FC in data/external_data/:
- 'SupplemenrayDataPfister.xlsx' (Appendix B, version 2.0 of 2013-08-07): sheet 'Blue water consumption (m3_ton)',
  production-weighted country values of the expected consumption (geometric mean of deficit and full irrigation,
  paper Eq. 1); sheet 'Global Averages WFP_mon and WSI', world production and expected consumption per crop;
- 'Water consumption of crop on watershed level (blue.zip': 'watershed_WF_crops.xlsx' (per intersection of a
  WaterGAP watershed and a country, FID: production, expected, minimum = deficit irrigation, maximum = full
  irrigation, m3/t) and 'BlueWGeoMean_Watershed_monthly.xlsx' (expected consumption per month, m3/t).

Outputs (crops of data/pfister2014_crop_mapping.csv, countries of data/pfister2014_country_mapping.csv):
- data/pfister2014_blue_water_country.csv: expected = Appendix B value (the published country value; empty where
  Appendix B says 'No production'); minimum, maximum and production = production-weighted over the FIDs of the
  country in the watershed file; GLO: expected and production of the sheet 'Global Averages', minimum and
  maximum production-weighted over all FIDs.
- data/pfister2014_blue_water_watershed.csv: one row per crop, tool country and FID with production: production,
  expected, minimum, maximum (m3/t), share of the country's blue water (expected x production / country sum) and
  the monthly profile (month value / sum of the 12 months of the monthly file).

Usage (repository root):
    python scripts/build_pfister_blue_water.py
    python scripts/build_pfister_blue_water.py --supplementary <xlsx> --watershed-zip <zip> --out-dir <dir>
"""
import argparse
import csv
import io
import sys
import warnings
import zipfile
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src" / "legacy_model"))

import dataloader  # noqa: E402

SUPPLEMENTARY = "SupplemenrayDataPfister.xlsx"
WATERSHED_ZIP = "Water consumption of crop on watershed level (blue.zip"
WATERSHED_FILE = "watershed_WF_crops.xlsx"
MONTHLY_FILE = "BlueWGeoMean_Watershed_monthly.xlsx"
MONTHS = ["Ja", "Fe", "Mr", "Ap", "My", "Jn", "Jl", "Au", "Se", "Oc", "No", "De"]
NO_PRODUCTION = "No production"
COUNTRY_COLUMNS = ["pfister_crop", "country", "bw_expected_m3_t", "bw_min_m3_t", "bw_max_m3_t", "production_t",
                   "n_watersheds", "note"]
WATERSHED_COLUMNS = ["pfister_crop", "country", "fid", "fipsbasin", "bas34s_id", "production_t", "bw_expected_m3_t",
                     "bw_min_m3_t", "bw_max_m3_t", "blue_water_share"] + ["month_%02d" % m for m in range(1, 13)]


def _num(value):
    return float(value) if isinstance(value, (int, float)) else 0.0


def _sig(value, digits=6):
    return float("%.*g" % (digits, value))


def _open_workbook(source):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(source, read_only=True, data_only=True)


def _zip_member(archive, filename):
    names = [i.filename for i in archive.infolist() if i.filename.endswith("/" + filename) or i.filename == filename]
    if len(names) != 1:
        raise ValueError("%s: expected one %s, found %s" % (archive.filename, filename, names))
    return _open_workbook(io.BytesIO(archive.read(names[0])))


def _mappings():
    crops = {}
    for row in dataloader.rows("pfister2014_crop_mapping.csv"):
        if row["use"] in ("direct", "proxy"):
            crops[row["pfister_crop"]] = row["pfister_crop_watershed"]
    countries = {row["country"]: row for row in dataloader.rows("pfister2014_country_mapping.csv")}
    return crops, countries


def read_supplementary(path):
    wb = _open_workbook(path)
    rows = list(wb["Blue water consumption (m3_ton)"].iter_rows(values_only=True))
    iso3 = list(rows[1])
    country_values = {row[0]: row for row in rows[2:] if row[0]}
    world = {}
    for row in wb["Global Averages WFP_mon and WSI"].iter_rows(values_only=True):
        if row[1] and isinstance(row[2], (int, float)) and isinstance(row[3], (int, float)):
            world[row[1]] = (float(row[2]), float(row[3]))
    return iso3, country_values, world


def read_watershed(archive):
    wb = _zip_member(archive, WATERSHED_FILE)
    sheets = {}
    for key, name in (("production", "Production_ton"), ("expected", "BW_per ton"), ("min", "BW_min_per ton"),
                      ("max", "BW_max_per ton")):
        rows = list(wb[name].iter_rows(values_only=True))
        sheets[key] = (list(rows[0]), [r for r in rows[1:] if r[0] is not None])
    header, fids = sheets["production"][0], [r[0] for r in sheets["production"][1]]
    for key, (h, body) in sheets.items():
        if h != header or [r[0] for r in body] != fids:
            raise ValueError("%s: sheet %s does not have the rows and columns of Production_ton" % (WATERSHED_FILE, key))
    return header, sheets


def read_monthly(archive):
    wb = _zip_member(archive, MONTHLY_FILE)
    out = {}
    for name in MONTHS:
        rows = list(wb[name].iter_rows(values_only=True))
        head = {r[1]: r for r in rows[:5]}
        out[name] = {r[1]: r for r in rows[5:] if r[1]}
    columns = {fid: j for j, fid in enumerate(head["FID_"]) if isinstance(fid, int)}
    return head, columns, out


def build(supplementary, watershed_zip):
    crops, countries = _mappings()
    iso3, country_values, world = read_supplementary(supplementary)
    with zipfile.ZipFile(watershed_zip) as archive:
        header, sheets = read_watershed(archive)
        head, month_columns, monthly = read_monthly(archive)
    name_to_country = {}
    for country, row in countries.items():
        for name in row["watershed_names"].split(";"):
            name_to_country[name] = country
    country_rows, watershed_rows = [], []
    for crop in sorted(crops):
        column = crops[crop]
        if header.count(column) != 1:
            raise ValueError("%s: column %r found %d times" % (WATERSHED_FILE, column, header.count(column)))
        j = header.index(column)
        if crop not in country_values or crop not in world:
            raise ValueError("%s: no row %r" % (SUPPLEMENTARY, crop))
        units = defaultdict(list)
        totals = [0.0, 0.0, 0.0]
        for i, prod_row in enumerate(sheets["production"][1]):
            production = _num(prod_row[j])
            if production <= 0:
                continue
            values = tuple(_num(sheets[k][1][i][j]) for k in ("expected", "min", "max"))
            totals[0] += production
            totals[1] += values[1] * production
            totals[2] += values[2] * production
            country = name_to_country.get(prod_row[1])
            if country is not None:
                units[country].append((prod_row[0], production) + values)
        for country in sorted(countries):
            value = country_values[crop][iso3.index(countries[country]["supplementary_iso3"])]
            ws = units.get(country, [])
            production = sum(u[1] for u in ws)
            if value == NO_PRODUCTION and not ws:
                continue
            notes = []
            if value == NO_PRODUCTION:
                notes.append("No production in Appendix B")
            if not ws:
                notes.append("no production in the watershed file")
            country_rows.append({
                "pfister_crop": crop, "country": country,
                "bw_expected_m3_t": "" if value == NO_PRODUCTION else float(value),
                "bw_min_m3_t": _sig(sum(u[3] * u[1] for u in ws) / production) if ws else "",
                "bw_max_m3_t": _sig(sum(u[4] * u[1] for u in ws) / production) if ws else "",
                "production_t": _sig(production) if ws else "", "n_watersheds": len(ws), "note": "; ".join(notes)})
            blue = sum(u[2] * u[1] for u in ws)
            for fid, prod, expected, low, high in sorted(ws):
                jm = month_columns[fid]
                months = [_num(monthly[m][crop][jm]) for m in MONTHS]
                total = sum(months)
                watershed_rows.append(dict(
                    {"pfister_crop": crop, "country": country, "fid": fid, "fipsbasin": head["FIPSBASIN"][jm],
                     "bas34s_id": head["BAS34S_ID"][jm], "production_t": _sig(prod), "bw_expected_m3_t": _sig(expected),
                     "bw_min_m3_t": _sig(low), "bw_max_m3_t": _sig(high),
                     "blue_water_share": _sig(expected * prod / blue) if blue > 0 else ""},
                    **{"month_%02d" % (k + 1): (_sig(v / total) if total > 0 else "") for k, v in enumerate(months)}))
        production, expected = world[crop]
        country_rows.append({"pfister_crop": crop, "country": "GLO", "bw_expected_m3_t": expected,
                             "bw_min_m3_t": _sig(totals[1] / totals[0]), "bw_max_m3_t": _sig(totals[2] / totals[0]),
                             "production_t": _sig(production), "n_watersheds": "",
                             "note": "expected and production: sheet Global Averages; minimum and maximum: all FIDs"})
    return country_rows, watershed_rows


def _write(path, columns, rows):
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print("%d rows written to %s" % (len(rows), path))


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--supplementary", default=str(dataloader.EXTERNAL_DATA_DIR / SUPPLEMENTARY))
    parser.add_argument("--watershed-zip", default=str(dataloader.EXTERNAL_DATA_DIR / WATERSHED_ZIP))
    parser.add_argument("--out-dir", default=str(dataloader.DATA_DIR))
    args = parser.parse_args()
    country_rows, watershed_rows = build(args.supplementary, args.watershed_zip)
    _write(Path(args.out_dir) / "pfister2014_blue_water_country.csv", COUNTRY_COLUMNS, country_rows)
    _write(Path(args.out_dir) / "pfister2014_blue_water_watershed.csv", WATERSHED_COLUMNS, watershed_rows)


if __name__ == "__main__":
    main()
