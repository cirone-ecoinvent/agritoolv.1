"""Build data/aquastat_irrigation_technique_area.csv, the area equipped for irrigation per technique and country.

Source (decision L44): FAO AQUASTAT, variables 'Area equipped for full control irrigation: surface
irrigation', '...: sprinkler irrigation' and '...: localized irrigation' (1000 ha), as compiled for the
ecoinvent irrigation activities (production volumes of 'irrigation, surface / sprinkler / drip') in the
workbook 'AQUASTAT Dissemination System LASTUP.xlsx' (export of the AQUASTAT Dissemination System of
2025-07-31, provided by FC):
- sheet 'Data': the AQUASTAT export (year 2020), countries matched by M49 code
  (data/faostat_country_mapping.csv); a row without M49 code (Colombia, surface irrigation, comment
  'Added manully') is matched by its ecoinvent geography and marked in the origin column;
- sheet 'CH_BE_RU_CA-QC_VN': only for the countries without any row in the export (BE 2005, CH 2006,
  RU 2006). Area = column 'right one' (m3) / application rate of the technique (sheet 'Irrigation
  rate', m3/ha/yr) / 1000, i.e. the area the workbook itself uses. CA and VN of this sheet are not used:
  the export has them (2020).
A technique that the workbook does not report for a country has no row.

Usage (repository root):
    python scripts/build_aquastat_irrigation_techniques.py                               # workbook in data/external_data/
    python scripts/build_aquastat_irrigation_techniques.py --workbook <xlsx> --out <path>
"""
import argparse
import csv
import sys
import warnings
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src" / "legacy_model"))

import dataloader  # noqa: E402

WORKBOOK = "AQUASTAT Dissemination System LASTUP.xlsx"
VARIABLE_PREFIX = "Area equipped for full control irrigation: "
TECHNIQUES = {"surface irrigation": "surface", "sprinkler irrigation": "sprinkler", "localized irrigation": "localized"}
RATE_ROWS = {"Surface irrigation": "surface", "Sprinkler irrigation": "sprinkler", "Drip irrigation": "localized"}
COLUMNS = ["country", "technique", "area_kha", "year", "aquastat_symbol", "origin"]


def _records(ws):
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    for values in rows:
        if any(v is not None for v in values):
            yield dict(zip(header, values))


def _technique(variable):
    if not variable.startswith(VARIABLE_PREFIX) or variable[len(VARIABLE_PREFIX):] not in TECHNIQUES:
        raise ValueError("unexpected AQUASTAT variable %r" % variable)
    return TECHNIQUES[variable[len(VARIABLE_PREFIX):]]


def build(workbook):
    countries = dataloader.rows("faostat_country_mapping.csv")
    m49_to_country = {int(r["m49_code"]): r["country"] for r in countries if r["country"] != "GLO"}
    tool = set(m49_to_country.values())
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    rates = {}
    for values in wb["Irrigation rate"].iter_rows(values_only=True):
        if values[0] in RATE_ROWS:
            rates[RATE_ROWS[values[0]]] = float(values[2])
    if set(rates) != set(TECHNIQUES.values()):
        raise ValueError("application rates not found in sheet 'Irrigation rate'")

    out = {}
    for d in _records(wb["Data"]):
        if d["m49"] is not None:
            country, origin = m49_to_country.get(int(d["m49"])), "AQUASTAT export (sheet Data)"
        else:
            country = d["ecogeo"] if d["ecogeo"] in tool else None
            origin = "sheet Data, row without M49 code, comment '%s'" % (d["Comment"] or "").strip()
        if country is None or d["Value"] is None:
            continue
        if d["Unit"] != "1000 ha":
            raise ValueError("unit %r, expected '1000 ha'" % d["Unit"])
        key = (country, _technique(d["Variable"]))
        if key in out:
            raise ValueError("two rows for %s %s" % key)
        out[key] = {"area_kha": float(d["Value"]), "year": d["Year"], "aquastat_symbol": d["Symbol"], "origin": origin}
    exported = {country for country, _ in out}
    for d in _records(wb["CH_BE_RU_CA-QC_VN"]):
        country = (d["ecogeo"] or "").strip()
        if country not in tool or country in exported:
            continue
        technique = _technique(d["Variable"])
        out[(country, technique)] = {"area_kha": float(d["right one"]) / rates[technique] / 1000.0, "year": d["Year"],
                                     "aquastat_symbol": "", "origin": "sheet CH_BE_RU_CA-QC_VN"}
    missing = sorted(tool - {country for country, _ in out})
    if missing:
        raise ValueError("no AQUASTAT area for %s" % ", ".join(missing))
    order = list(TECHNIQUES.values())
    return [dict(country=country, technique=technique, area_kha=round(v["area_kha"], 6), year=v["year"],
                 aquastat_symbol=v["aquastat_symbol"], origin=v["origin"])
            for (country, technique), v in sorted(out.items(), key=lambda kv: (kv[0][0], order.index(kv[0][1])))]


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", default=str(dataloader.EXTERNAL_DATA_DIR / WORKBOOK),
                        help="AQUASTAT workbook (default: data/external_data/%s)" % WORKBOOK)
    parser.add_argument("--out", default=str(dataloader.DATA_DIR / "aquastat_irrigation_technique_area.csv"))
    args = parser.parse_args()
    rows = build(args.workbook)
    with open(args.out, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print("%d rows written to %s" % (len(rows), args.out))


if __name__ == "__main__":
    main()
