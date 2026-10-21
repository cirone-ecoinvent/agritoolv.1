"""Build data/atlas2026_pesticides_country.csv, the total pesticide active ingredients (a.i.) applied per ha of the
tool crops and countries according to the Global atlas of pesticide use, decision L48.

Sources:
  - Zhang Y., Gladbach A., Kosnik M., Settles L., Norinder U., Glas D., Fantke P. (2026) Global atlas of pesticide use.
    Research Square preprint, version 1, doi 10.21203/rs.3.rs-10665000/v1, CC BY 4.0:
      Supplementary data file 1, Table S1 (annual average applied mass 2016-2020 per country or region, crop and
      active ingredient), data/external_data/'Global atlas of pesticide use 1.xlsx';
      Supplementary data file 2, Table S4 (country names -> ISO 3166-1 alpha-3) and Table S6 (crop names -> FAOSTAT
      crop code), data/external_data/'Global atlas of pesticide use.xlsx';
  - FAOSTAT QCL, element 5312 Area harvested (ha), data/external_data/Production_Crops_Livestock_E_All_Data_
    (Normalized).csv (the download of the default yield, L23);
  - UNSD, Standard country or area codes for statistical use (M49), data/external_data/'UNSD — Methodology.csv'
    (ISO alpha-3 -> M49).

Method:
  - atlas crops: the crops of Table S6 whose FAOSTAT crop code is the item of a tool crop in
    faostat_crop_item_mapping.csv; countries: Table S1 country name -> ISO alpha-3 (Table S4) -> M49 (UNSD) -> tool
    country (faostat_country_mapping.csv);
  - applied mass = sum of 'Annual applied mass [kg]' over all rows of the crop and country: all a.i., target classes
    and treatments (crop, seed, soil), soil fumigants included (FC, 2026-10-21); soil_fumigant_mass_kg = the part of
    the a.i. of soil_fumigant_active_ingredients.csv (information only);
  - area = mean FAOSTAT area harvested 2016-2020 of the item (years without a value skipped), so the value is per ha
    and crop cycle as the yield (L24); the annual load of the atlas (per ha of physical area, mapSPAM 2020) is not in
    the public files;
  - kg_ai_per_ha = applied mass / area; GLO = applied mass of all countries and regions of Table S1 / FAOSTAT World.
Flags (decision L48): a tool-country value is flagged when it has no FAOSTAT area, at most
pesticide_min_active_ingredients a.i., or lies outside a factor pesticide_outlier_factor of the median of the
tool-country values of the crop (model_parameters.csv). The use of every flagged value (atlas or fallback) is decided
in atlas2026_pesticides_review.csv; flagged values missing there are printed.

Usage (repository root):
    python scripts/build_atlas_pesticides.py
    python scripts/build_atlas_pesticides.py --s1 <xlsx> --s2 <xlsx> --qcl <csv> --m49 <csv> --out <path>
"""
import argparse
import csv
import statistics
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src" / "legacy_model"))

import dataloader  # noqa: E402

S1_FILE = "Global atlas of pesticide use 1.xlsx"
S2_FILE = "Global atlas of pesticide use.xlsx"
QCL_FILE = "Production_Crops_Livestock_E_All_Data_(Normalized).csv"
M49_FILE = "UNSD — Methodology.csv"
S1_COLUMNS = {"country": "Country/region name", "crop": "Crop name", "cas": "CAS RN", "mass": "Annual applied mass [kg]"}
AREA_HARVESTED = ("5312", "ha")
AREA_YEARS = range(2016, 2021)
WORLD = "GLO"
COLUMNS = ["atlas_crop", "faostat_item_code", "country", "n_active_ingredients", "applied_mass_kg",
           "soil_fumigant_mass_kg", "area_harvested_ha", "area_years", "kg_ai_per_ha"]
REVIEW = "atlas2026_pesticides_review.csv"


def _round(value, digits=6):
    return float("%.*g" % (digits, value))


def read_s2(path):
    """({atlas crop: FAOSTAT crop code}, {country name: ISO alpha-3}) from Tables S6 and S4."""
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True)
    rows = workbook["Table S6"].iter_rows(values_only=True)
    index = {name: i for i, name in enumerate(next(rows))}
    items = {}
    for row in rows:
        crop, code = row[index["Crop, harmonized"]], row[index["Crop code, FAOSTAT"]]
        if crop and code not in (None, "#N/A"):
            if items.setdefault(crop, str(code)) != str(code):
                raise ValueError("Table S6: %s has two FAOSTAT codes" % crop)
    rows = workbook["Table S4"].iter_rows(values_only=True)
    index = {name: i for i, name in enumerate(next(rows))}
    iso3 = {row[index["Country/region, harmonized"]].strip(): row[index["ISO 3166-1 alpha-3"]].strip()
            for row in rows if row[index["Country/region, harmonized"]]}
    return items, iso3


def read_s1(path, crops, country_of_name, fumigants):
    """{(crop, country): [mass, fumigant mass, set of CAS]} for the given crops and the tool countries, and
    {crop: mass} over all countries and regions (country WORLD)."""
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True)
    rows = workbook["Table S1"].iter_rows(values_only=True)
    next(rows)  # group headers
    index = {name: i for i, name in enumerate(next(rows))}
    col = {key: index[name] for key, name in S1_COLUMNS.items()}
    data = defaultdict(lambda: [0.0, 0.0, set()])
    for row in rows:
        if row[col["crop"]] is None:
            continue
        crop = row[col["crop"]].strip().lower()
        if crop not in crops:
            continue
        mass, cas = float(row[col["mass"]] or 0.0), row[col["cas"]]
        for country in (WORLD, country_of_name.get(row[col["country"]].strip())):
            if country:
                entry = data[(crop, country)]
                entry[0] += mass
                entry[1] += mass if cas in fumigants else 0.0
                entry[2].add(cas)
    return data


def read_area(path, items, areas):
    """{(item, area code): (mean ha over AREA_YEARS, number of years)} of the QCL element Area harvested."""
    values = defaultdict(list)
    with open(path, encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            if (row["Element Code"] != AREA_HARVESTED[0] or row["Item Code"] not in items
                    or row["Area Code"] not in areas or int(row["Year"]) not in AREA_YEARS or not row["Value"].strip()):
                continue
            if row["Unit"] != AREA_HARVESTED[1]:
                raise ValueError("QCL area unit %r, expected %r" % (row["Unit"], AREA_HARVESTED[1]))
            values[(row["Item Code"], row["Area Code"])].append(float(row["Value"]))
    return {key: (statistics.mean(v), len(v)) for key, v in values.items() if v}


def read_m49(path):
    """{ISO alpha-3: M49 code (3 digits)} from the UNSD overview table (semicolon separated)."""
    with open(path, encoding="utf-8-sig", newline="") as f:
        return {row["ISO-alpha3 Code"]: row["M49 Code"].zfill(3) for row in csv.DictReader(f, delimiter=";")
                if row["ISO-alpha3 Code"]}


def build(s1, s2, qcl, m49_file):
    items, iso3 = read_s2(s2)
    tool_items = {row["faostat_item_code"] for row in dataloader.rows("faostat_crop_item_mapping.csv")}
    crops = {crop: item for crop, item in items.items() if item in tool_items}
    countries = {row["m49_code"]: row for row in dataloader.rows("faostat_country_mapping.csv")}
    m49 = read_m49(m49_file)
    country_of_name = {name: countries[m49[code]]["country"] for name, code in iso3.items()
                       if code in m49 and m49[code] in countries and countries[m49[code]]["country"] != WORLD}
    fumigants = {row["cas_rn"] for row in dataloader.rows("soil_fumigant_active_ingredients.csv")}
    data = read_s1(s1, crops, country_of_name, fumigants)
    area_code = {row["country"]: row["faostat_area_code"] for row in countries.values()}
    area = read_area(qcl, set(crops.values()), set(area_code.values()))
    out = []
    for (crop, country), (mass, fumigant, cas) in data.items():
        ha, years = area.get((crops[crop], area_code[country]), (None, 0))
        out.append({"atlas_crop": crop, "faostat_item_code": crops[crop], "country": country,
                    "n_active_ingredients": len(cas), "applied_mass_kg": _round(mass),
                    "soil_fumigant_mass_kg": _round(fumigant), "area_harvested_ha": _round(ha) if ha else "",
                    "area_years": years, "kg_ai_per_ha": _round(mass / ha) if ha else ""})
    out.sort(key=lambda row: (row["atlas_crop"], row["country"] != WORLD, row["country"]))
    return out


def flags(rows, factor=None, min_ai=None):
    """{(atlas crop, country): flag} of the tool-country values that need a decision (see the module docstring)."""
    factor = factor or dataloader.param("pesticide_outlier_factor")
    min_ai = min_ai or dataloader.param("pesticide_min_active_ingredients")
    values = defaultdict(dict)
    for row in rows:
        if row["country"] != WORLD and row["kg_ai_per_ha"] != "":
            values[row["atlas_crop"]][row["country"]] = float(row["kg_ai_per_ha"])
    median = {crop: statistics.median(v.values()) for crop, v in values.items()}
    out = {}
    for row in rows:
        crop, country = row["atlas_crop"], row["country"]
        if country == WORLD:
            continue
        if row["kg_ai_per_ha"] == "":
            out[(crop, country)] = "no_area"
        elif int(row["n_active_ingredients"]) <= min_ai:
            out[(crop, country)] = "few_active_ingredients"
        elif float(row["kg_ai_per_ha"]) < median[crop] / factor:
            out[(crop, country)] = "below_median_range"
        elif float(row["kg_ai_per_ha"]) > median[crop] * factor:
            out[(crop, country)] = "above_median_range"
    return out


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--s1", default=str(dataloader.EXTERNAL_DATA_DIR / S1_FILE))
    parser.add_argument("--s2", default=str(dataloader.EXTERNAL_DATA_DIR / S2_FILE))
    parser.add_argument("--qcl", default=str(dataloader.EXTERNAL_DATA_DIR / QCL_FILE))
    parser.add_argument("--m49", default=str(dataloader.EXTERNAL_DATA_DIR / M49_FILE))
    parser.add_argument("--out", default=str(dataloader.DATA_DIR / "atlas2026_pesticides_country.csv"))
    args = parser.parse_args()
    rows = build(args.s1, args.s2, args.qcl, args.m49)
    with open(args.out, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print("%d rows written to %s" % (len(rows), args.out))
    flagged = flags([{key: str(value) for key, value in row.items()} for row in rows])
    reviewed = {(row["atlas_crop"], row["country"]) for row in dataloader.rows(REVIEW)} \
        if (dataloader.DATA_DIR / REVIEW).is_file() else set()
    print("%d flagged values, not in %s: %s" % (len(flagged), REVIEW,
                                               sorted(set(flagged) - reviewed) or "none"))


if __name__ == "__main__":
    main()
