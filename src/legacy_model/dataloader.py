"""Single data-access module for all external parameters and lookup tables.

Every constant used by the models (emission factors, model parameters, heavy-metal contents,
mappings, ecoinvent master data) is read from CSV files in ``<repo>/data`` through this module,
so sources can be updated or replaced without touching the calculation code. The provenance of
each file is recorded in ``data/sources.csv`` and ``data/SOURCES.md``.

CSV conventions: UTF-8, comma delimiter, dot decimal separator, one header row, units stated in
the column names or in a dedicated ``unit`` column. Empty cells mean "not available" (NA).
"""
import csv
import xml.etree.ElementTree as ET
from collections import namedtuple
from functools import lru_cache
from pathlib import Path

DATA_DIR = Path(__file__).resolve().parents[2] / "data"

IntermediateExchange = namedtuple("IntermediateExchange", ["uuid", "name", "unit"])


def data_path(filename):
    """Absolute path of a data file; raises if it does not exist."""
    path = DATA_DIR / filename
    if not path.is_file():
        raise FileNotFoundError("Data file not found: %s" % path)
    return path


@lru_cache(maxsize=None)
def rows(filename):
    """All rows of a CSV file as a tuple of dicts (values as strings).

    Completely empty rows (e.g. ",,,,," lines appended by spreadsheet-like editors) are skipped.
    """
    with open(data_path(filename), newline="", encoding="utf-8") as f:
        return tuple(row for row in csv.DictReader(f)
                     if any((value or "").strip() for value in row.values()))


@lru_cache(maxsize=None)
def params():
    """Scalar model parameters from ``model_parameters.csv`` as {name: float}."""
    return {row["name"]: float(row["value"]) for row in rows("model_parameters.csv")}


def param(name):
    """One scalar parameter from ``model_parameters.csv``."""
    try:
        return params()[name]
    except KeyError:
        raise KeyError("Parameter '%s' not found in model_parameters.csv" % name)


def _to_float_or_none(value):
    return float(value) if value not in ("", None) else None


@lru_cache(maxsize=None)
def table(filename, key_column=None, numeric=True):
    """A CSV table keyed by its first column (or ``key_column``).

    Returns {key: {column: value}}. With ``numeric=True`` every non-key column is converted to
    float, empty cells become None; columns that cannot be converted are kept as strings.
    """
    out = {}
    for row in rows(filename):
        key = row[key_column] if key_column else row[next(iter(row))]
        rec = {}
        for col, val in row.items():
            if col == (key_column or next(iter(row))):
                continue
            if numeric:
                try:
                    rec[col] = _to_float_or_none(val)
                except ValueError:
                    rec[col] = val
            else:
                rec[col] = val
        out[key] = rec
    return out


def metal_row(filename, key, metals):
    """Heavy-metal vector {metal: mg value} for one table row; NA (empty cell) -> 0.0."""
    row = table(filename)[key]
    return {m: (row.get(m) or 0.0) for m in metals}


def single_metal_row(filename, metals):
    """Heavy-metal vector of a one-row table (deposition, leaching); NA -> 0.0."""
    row = next(iter(table(filename).values()))
    return {m: (row.get(m) or 0.0) for m in metals}


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def load_intermediate_exchanges(path, sheet="IntermediateExchanges", name_col="name",
                                id_col="id", unit_col="unitName"):
    """Load ecoinvent Master Data intermediate exchanges, keyed by exact English name.

    Accepts the ecoSpold2 master-data XML (``IntermediateExchanges.xml``, namespace-agnostic) or
    a table (``.xlsx`` sheet / ``.csv``) with the given column names. Returns
    {name: IntermediateExchange(uuid, name, unit)}.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError("Master data file not found: %s" % path)
    records = []
    suffix = path.suffix.lower()
    if suffix == ".xml":
        for el in ET.parse(path).getroot().iter():
            if _local(el.tag) != "intermediateExchange":
                continue
            name = unit = ""
            for child in el:
                lang = child.get("{http://www.w3.org/XML/1998/namespace}lang", "en")
                if _local(child.tag) == "name" and lang == "en":
                    name = (child.text or "").strip()
                elif _local(child.tag) == "unitName" and lang == "en":
                    unit = (child.text or "").strip()
            records.append((el.get("id", ""), name, unit))
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook  # optional dependency, only for xlsx master data
        ws = load_workbook(path, read_only=True, data_only=True)[sheet]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        try:
            i_id, i_name, i_unit = (header.index(c) for c in (id_col, name_col, unit_col))
        except ValueError:
            raise ValueError("Sheet '%s' must have columns %r; found %r"
                             % (sheet, (id_col, name_col, unit_col), header))
        records = [(str(r[i_id]), str(r[i_name]).strip(), str(r[i_unit]))
                   for r in it if r[i_name]]
    else:
        with open(path, newline="", encoding="utf-8") as f:
            records = [(r[id_col], r[name_col].strip(), r[unit_col]) for r in csv.DictReader(f)]
    return {name: IntermediateExchange(uuid, name, unit) for uuid, name, unit in records if name}
