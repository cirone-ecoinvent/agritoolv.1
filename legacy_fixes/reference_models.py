"""Corrected reference implementations for the legacy sri-crop-tool field-emission models.

Specification: Quantis (2019), "Models integrated in ecoinvent LCI calculation tool for
crop production". Where that document is silent or wrong, the choice is backed by:
* EEA (2016) EMEP/EEA Guidebook, chapter 3.D (NOx factor);
* Nemecek et al. (2023) SALCA, Int J LCA, Online Resources ESM2 (SALCAnitrate),
  ESM5 (SALCAfieldP) and ESM7 (SALCAheavymetal).
Section numbers in docstrings refer to the Quantis document. Parameters come from data/*.csv.

Conventions
-----------
* Quantities per hectare and per inventory period. ``duration_yr`` = occupation period
  in years (harvest of previous crop -> harvest of this crop; 1.0 if unknown).
* ``drained_fraction`` = share of the field with drainage (default 0 = document behaviour,
  everything to ground water). Used consistently for NO3, phosphate and heavy metals.
* Gases returned as emitted species (kg NH3, NO2, N2O, NO3, PO4); N bookkeeping in kg N.
* Heavy metals handled in mg internally, returned in kg.

Nitrogen calculation order: NH3 -> NOx -> direct N2O -> NO3 leaching -> indirect N2O.
"""
from __future__ import annotations

import csv
import difflib
import re
import warnings
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Iterable, Mapping

DATA_DIR = Path(__file__).parent / "data"
METALS = ("Cd", "Cu", "Zn", "Pb", "Ni", "Cr", "Hg")

# Stoichiometric conversion factors (molar-mass ratios, not external data)
N_TO_NH3 = 17 / 14
N_TO_NO2 = 46 / 14
N_TO_NO3 = 62 / 14
N_TO_N2O = 44 / 28
P_TO_PO4 = 94.971 / 30.974
MG_TO_KG = 1e-6


# --------------------------------------------------------------------------- data access
@lru_cache(maxsize=None)
def params() -> dict[str, float]:
    """Scalar model parameters from data/model_parameters.csv."""
    with open(DATA_DIR / "model_parameters.csv", newline="", encoding="utf-8") as f:
        return {row["name"]: float(row["value"]) for row in csv.DictReader(f)}


def _p(name: str) -> float:
    return params()[name]


@lru_cache(maxsize=None)
def _metal_table(filename: str) -> dict[str, dict[str, float | str | None]]:
    """Heavy-metal table keyed by its first column; empty cells -> None (NA)."""
    out: dict[str, dict[str, float | str | None]] = {}
    with open(DATA_DIR / filename, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        key_col = reader.fieldnames[0]
        for row in reader:
            rec: dict[str, float | str | None] = {}
            for col, val in row.items():
                if col == key_col:
                    continue
                if col == "nutrient":
                    rec[col] = val
                else:
                    rec[col] = float(val) if val not in ("", None) else None
            out[row[key_col]] = rec
    return out


def _single_row(filename: str) -> dict[str, float]:
    row = next(iter(_metal_table(filename).values()))
    return {m: (row.get(m) or 0.0) for m in METALS}  # NA -> 0 (Ni leaching)


def _check_fraction(name: str, value: float) -> None:
    if not 0.0 <= value <= 1.0:
        raise ValueError(f"{name} must be in [0, 1], got {value}")


# --------------------------------------------------------------------------- NH3 (2.1)
def nh3_from_nh3_n(nh3_n_kg: float) -> float:
    """Convert NH3-N [kg N] to kg NH3 (17/14). Applies to every NH3 path, incl. compost
    and sewage sludge."""
    return nh3_n_kg * N_TO_NH3


def nh3_mineral(applications: Iterable[tuple[float, float, float]], p_acidic: float) -> float:
    """NH3 from mineral fertilisers, doc 2.1: 17/14 * sum_m (EFa_m*p + EFb_m*(1-p)) * N_m.

    ``applications``: (n_applied_kg, efa, efb) per fertiliser type, EF in kg NH3-N/kg N.
    """
    nh3_n = sum(n * (efa * p_acidic + efb * (1 - p_acidic)) for n, efa, efb in applications)
    return nh3_from_nh3_n(nh3_n)


# --------------------------------------------------------------------------- NOx (2.3)
def nox(n_applied_total: float, nh3_kg: float) -> float:
    """NOx [kg NO2] from mineral + organic N.

    EF = 0.04 kg NO2/kg N (EEA 2016 3.D Tab. 3.1, note a: reported as NO2), i.e.
    0.012174 kg NOx-N/kg N. Deviation from Quantis 2019, which converted with 14/30 as if
    the factor were in NO (+53 %). Subtraction of NH3-N kept as in Quantis/WFLDB.
    """
    nh3_n = nh3_kg / N_TO_NH3
    n_base = max(n_applied_total - nh3_n, 0.0)
    return _p("ef_nox_n") * n_base * N_TO_NO2


# --------------------------------------------------------------------------- N2O (2.2)
def n_som_from_soc_change(delta_soc_kg_c: float, cn_ratio: float) -> float:
    """N mineralised from SOC loss (IPCC F_SOM) [kg N]; only losses count.

    ``delta_soc_kg_c`` from the LUC / land-management module (negative = loss);
    ``cn_ratio`` 15 (forest/grassland -> cropland) or 11 (cropland), doc 2.9.1.2.
    Not the SQCB soil N_org. As in SALCAfieldN (Nemecek et al. 2023, ESM1 Eq. 5), the N2O
    from SOC mineralisation is reported only here, in the field N2O; the LUC module
    supplies the SOC change and must not report this N2O itself.
    """
    return max(-delta_soc_kg_c, 0.0) / cn_ratio


def n2o_direct_n(n_applied_total: float, n_crop_residues: float = 0.0,
                 n_som: float = 0.0, flooded_rice: bool = False) -> float:
    """Direct N2O-N [kg N], IPCC 2006 Eq. 11.1 with EF1 or EF1FR (flooded rice)."""
    ef1 = _p("ef1_flooded_rice") if flooded_rice else _p("ef1_direct_n2o")
    return ef1 * (n_applied_total + n_crop_residues + n_som)


@dataclass(frozen=True)
class N2OResult:
    direct_n: float
    indirect_volatilisation_n: float
    indirect_leaching_n: float

    @property
    def total_n(self) -> float:
        return self.direct_n + self.indirect_volatilisation_n + self.indirect_leaching_n

    @property
    def total_n2o(self) -> float:
        return self.total_n * N_TO_N2O


def n2o(n_applied_total: float, n_crop_residues: float, n_som: float,
        nh3_kg: float, nox_kg: float, no3_kg: float, flooded_rice: bool = False) -> N2OResult:
    """N2O, doc 2.2 = IPCC 2006 Eq. 11.1 + 11.9 + 11.10 with modelled NH3, NOx, NO3.

    EF1FR applies to the direct term only; EF4 and EF5 are unchanged for rice.
    ``no3_kg`` is total nitrate leached (ground + surface water).
    """
    direct = n2o_direct_n(n_applied_total, n_crop_residues, n_som, flooded_rice)
    vol = _p("ef4_volatilisation") * (nh3_kg / N_TO_NH3 + nox_kg / N_TO_NO2)
    leach = _p("ef5_leaching") * (no3_kg / N_TO_NO3)
    return N2OResult(direct, vol, leach)


# --------------------------------------------------------------------------- NO3 (2.4)
def soil_organic_n(soil_c_fraction: float) -> float:
    """Organic N stock in the upper 50 cm [kg N/ha], doc 2.4 (Annex 1 ratio, e.g. 0.011)."""
    if not 0.0 < soil_c_fraction < 1.0:
        raise ValueError(f"soil_c_fraction must be a ratio (Annex 1), got {soil_c_fraction}")
    n_tot = soil_c_fraction * _p("soil_volume_m3") * _p("bulk_density_kg_m3") / _p("soil_cn_ratio")
    return n_tot * _p("norg_ntot_ratio")


@dataclass(frozen=True)
class NitrateResult:
    n_supply_net: float      # S [kg N]
    no3_n: float             # total leached NO3-N [kg N]
    drained_fraction: float = 0.0

    @property
    def no3(self) -> float:
        return self.no3_n * N_TO_NO3

    @property
    def no3_groundwater(self) -> float:
        return self.no3 * (1 - self.drained_fraction)

    @property
    def no3_surface_water(self) -> float:
        return self.no3 * self.drained_fraction


def nitrate_leaching(*, precipitation_mm_yr: float, irrigation_mm_yr: float,
                     clay_fraction: float, rooting_depth_m: float,
                     n_mineral: float, n_organic_soluble: float,
                     nh3_n: float, nox_n: float, n2o_n: float,
                     n_uptake: float, soil_c_fraction: float,
                     duration_yr: float = 1.0, legume: bool = False,
                     drained_fraction: float = 0.0) -> NitrateResult:
    """SQCB-NO3 leaching, doc 2.4:
    N = 21.37 + P/(c*L) * [0.0037*S + 0.0000601*Norg - 0.00362*U]   [kg N/(ha*yr)]

    * S = mineral N + soluble organic N - (NH3-N + NOx-N + N2O-N); ``n2o_n`` = direct N2O-N
      from applied N (SALCAnitrate ESM2 also subtracts N2O losses from fertiliser N).
    * Per cycle: intercept and N_org term scaled by duration_yr, S and U as cycle totals
      (identical to the doc for t = 1; accounting period harvest-to-harvest as in ESM2).
    * c = clay in % (Annex 1 ratio x 100). P = annual precipitation + irrigation.
    * Split ground/surface water by drained fraction as in ESM2 Eq. 11-12.
    """
    if not 0.0 < clay_fraction < 1.0:
        raise ValueError(f"clay_fraction must be a ratio (Annex 1), got {clay_fraction}")
    if duration_yr <= 0:
        raise ValueError("duration_yr must be > 0")
    _check_fraction("drained_fraction", drained_fraction)
    t = duration_yr
    uptake = n_uptake * (_p("legume_uptake_share") if legume else 1.0)
    s_net = max(n_mineral + n_organic_soluble - nh3_n - nox_n - n2o_n, 0.0)
    k = (precipitation_mm_yr + irrigation_mm_yr) / (clay_fraction * 100.0 * rooting_depth_m)
    n_leach = _p("sqcb_intercept") * t + k * (
        _p("sqcb_coef_s") * s_net
        + _p("sqcb_coef_norg") * soil_organic_n(soil_c_fraction) * t
        - _p("sqcb_coef_u") * uptake
    )
    if n_leach < 0:
        warnings.warn(f"SQCB-NO3 regression negative ({n_leach:.2f} kg N); set to 0")
        n_leach = 0.0
    return NitrateResult(s_net, n_leach, drained_fraction)


# --------------------------------------------------------------------------- N orchestration
def nitrogen_emissions(*, n_mineral: float, n_organic_total: float, n_organic_soluble: float,
                       nh3_kg: float, n_crop_residues: float, n_som: float,
                       precipitation_mm_yr: float, irrigation_mm_yr: float,
                       clay_fraction: float, rooting_depth_m: float, n_uptake: float,
                       soil_c_fraction: float, duration_yr: float = 1.0,
                       legume: bool = False, flooded_rice: bool = False,
                       drained_fraction: float = 0.0) -> dict[str, float]:
    """All N field emissions in dependency order."""
    n_applied = n_mineral + n_organic_total
    nox_kg = nox(n_applied, nh3_kg)
    nitrate = nitrate_leaching(
        precipitation_mm_yr=precipitation_mm_yr, irrigation_mm_yr=irrigation_mm_yr,
        clay_fraction=clay_fraction, rooting_depth_m=rooting_depth_m,
        n_mineral=n_mineral, n_organic_soluble=n_organic_soluble,
        nh3_n=nh3_kg / N_TO_NH3, nox_n=nox_kg / N_TO_NO2,
        n2o_n=n2o_direct_n(n_applied, flooded_rice=flooded_rice),
        n_uptake=n_uptake, soil_c_fraction=soil_c_fraction,
        duration_yr=duration_yr, legume=legume, drained_fraction=drained_fraction,
    )
    n2o_res = n2o(n_applied, n_crop_residues, n_som, nh3_kg, nox_kg, nitrate.no3, flooded_rice)
    return {"NH3": nh3_kg, "NOx_as_NO2": nox_kg, "N2O": n2o_res.total_n2o,
            "NO3_groundwater": nitrate.no3_groundwater,
            "NO3_surface_water": nitrate.no3_surface_water}


# --------------------------------------------------------------------------- phosphorus (2.5)
_GW_KEYS = {"arable": "p_gw_arable", "grassland_intensive": "p_gw_grassland",
            "grassland_extensive": "p_gw_grassland"}
_RO_KEYS = {"arable": "p_ro_arable", "grassland_intensive": "p_ro_grassland_intensive",
            "grassland_extensive": "p_ro_grassland_extensive"}


@dataclass(frozen=True)
class PhosphateLeaching:
    groundwater_po4: float
    drainage_surface_water_po4: float


def phosphate_leaching(land_use: str, p2o5_slurry: float, duration_yr: float = 1.0,
                       drained_fraction: float = 0.0) -> PhosphateLeaching:
    """Phosphate leaching [kg PO4], doc 2.5.1 + drainage (SALCAfieldP ESM5 Eq. 4-5).

    Pg = Pgwl * (1 + 0.2/80 * P2O5sl); drained share: Pd = Pg * kd (kd = 6, Prasuhn 2006)
    to surface water. ASSUMPTION: on the drained share leaching goes to drains instead of
    ground water (same split as ESM2 Eq. 11-12 and ESM7 Eq. 5-6; ESM5 does not state it).
    P2O5 in kg as in Quantis/WFLDB (ESM5 states kg P: to verify for the new version).
    """
    _check_fraction("drained_fraction", drained_fraction)
    p_g = _p(_GW_KEYS[land_use]) * duration_yr * (1 + 0.2 / 80 * p2o5_slurry)
    return PhosphateLeaching(
        groundwater_po4=p_g * (1 - drained_fraction) * P_TO_PO4,
        drainage_surface_water_po4=p_g * _p("p_drainage_factor") * drained_fraction * P_TO_PO4,
    )


def phosphate_runoff(land_use: str, p2o5_mineral: float, p2o5_slurry: float,
                     p2o5_manure: float, duration_yr: float = 1.0,
                     slope_pct: float | None = None) -> float:
    """Phosphate run-off to surface water [kg PO4], doc 2.5.2 + slope factor (ESM5 Eq. 2).

    s = 0 for slope <= 3 %, 1 above (Prasuhn 2006); slope_pct=None -> s = 1 (doc default).
    Fro = 1 + 0.2/80*P2O5min + 0.7/80*P2O5sl + 0.4/80*P2O5man.
    """
    if slope_pct is not None and slope_pct <= _p("p_runoff_slope_threshold_pct"):
        return 0.0
    f_ro = 1 + 0.2 / 80 * p2o5_mineral + 0.7 / 80 * p2o5_slurry + 0.4 / 80 * p2o5_manure
    return _p(_RO_KEYS[land_use]) * duration_yr * f_ro * P_TO_PO4


def phosphorus_erosion(soil_eroded_kg: float) -> float:
    """Phosphorus to surface water by erosion [kg P], doc 2.5.3: Ser*Pcs*Fr*Ferw.

    ``soil_eroded_kg``: soil loss in the inventory period.
    """
    return (soil_eroded_kg * _p("p_topsoil_content") * _p("p_enrichment_factor")
            * _p("erosion_delivery_ratio"))


# --------------------------------------------------------------------------- heavy metals (2.7)
def _zero() -> dict[str, float]:
    return {m: 0.0 for m in METALS}


def hm_input_mineral(fertilisers: Mapping[str, float]) -> dict[str, float]:
    """Heavy-metal input [mg] from mineral fertilisers (Tab. 6), keyed by metal name.

    ``fertilisers``: {table key: kg of the row's reference nutrient (N, P2O5, K2O, CaO)}.
    Hg: no data -> 0.
    """
    table = _metal_table("hm_mineral_fertiliser_mg_per_kg_nutrient.csv")
    out = _zero()
    for name, kg_nutrient in fertilisers.items():
        row = table[name]
        for m in METALS:
            out[m] += kg_nutrient * (row.get(m) or 0.0)
    return out


def hm_input_compound(n_kg: float = 0.0, p2o5_kg: float = 0.0, k2o_kg: float = 0.0) -> dict[str, float]:
    """'Other' / compound fertilisers: each nutrient x generic-mean row of Tab. 6 [mg]."""
    return hm_input_mineral({"generic_mean_n": n_kg, "generic_mean_p": p2o5_kg,
                             "generic_mean_k": k2o_kg})


def hm_input_manure(manures: Mapping[str, float]) -> dict[str, float]:
    """Heavy-metal input [mg] from farmyard manure (Tab. 7): kg FM * DM * mg/kg DM."""
    table = _metal_table("hm_manure_mg_per_kg_dm.csv")
    out = _zero()
    for name, kg_fm in manures.items():
        row = table[name]
        kg_dm = kg_fm * row["dm_fraction"]
        for m in METALS:
            out[m] += kg_dm * (row.get(m) or 0.0)
    return out


def hm_plant_content(product: str) -> dict[str, float]:
    """Heavy-metal content [mg/kg DM] of a plant product (Tab. 5, Freiermuth 2006).

    NA -> 0, as in SALCAheavymetal 2023 (ESM7 Tab. 1). Products not in Tab. 5 -> generic
    mean, with a warning.
    """
    table = _metal_table("hm_plant_content_mg_per_kg_dm.csv")
    if product not in table:
        warnings.warn(f"'{product}' not in Tab. 5, generic mean used")
        product = "generic_mean"
    row = table[product]
    return {m: (row.get(m) or 0.0) for m in METALS}


def hm_biomass(products_kg_dm: Mapping[str, float]) -> dict[str, float]:
    """Heavy metals [mg] in biomass (exported products or seed), {product: kg DM}."""
    out = _zero()
    for product, kg_dm in products_kg_dm.items():
        content = hm_plant_content(product)
        for m in METALS:
            out[m] += kg_dm * content[m]
    return out


@dataclass(frozen=True)
class HeavyMetalResult:
    allocation: float
    leaching_groundwater_kg: float
    leaching_surface_water_kg: float
    erosion_kg: float
    soil_kg: float


def heavy_metals(*, agro_inputs_mg: Mapping[str, float], biomass_export_mg: Mapping[str, float],
                 soil_eroded_kg: float, land_use: str = "arable_land",
                 duration_yr: float = 1.0, drained_fraction: float = 0.0,
                 include_uptake: bool = True) -> dict[str, HeavyMetalResult]:
    """Heavy-metal balance, doc 2.7 written as SALCAheavymetal 2023 (ESM7 Eq. 3-7).

    A_i     = M_agro / (M_agro + m_dep * t)                 (0 if M_agro = 0)
    harvest = export * A_i;  leaching = m_leach * t * A_i (split by drained fraction)
    erosion = c_soil * Ser * 1.86 * 0.2 * A_i
    soil    = M_agro - harvest - leaching - erosion
    Identical to Quantis/Freiermuth Msoil = (inputs + deposition - unallocated outputs) * A_i:
    A_i is applied exactly once to each output.
    agro_inputs_mg: fertilisers + manure + seed + pesticides [mg], deposition excluded.
    include_uptake=False reproduces the "heavy_metal_uptake" switch.
    """
    _check_fraction("drained_fraction", drained_fraction)
    soil = _metal_table("hm_soil_content_mg_per_kg.csv")[land_use]
    dep = _single_row("hm_deposition_mg_per_ha_yr.csv")
    leach = _single_row("hm_leaching_mg_per_ha_yr.csv")
    a, f_er = _p("hm_accumulation_factor"), _p("hm_erosion_factor")
    out = {}
    for m in METALS:
        m_agro = agro_inputs_mg.get(m, 0.0)
        alloc = m_agro / (m_agro + dep[m] * duration_yr) if m_agro > 0 else 0.0
        harvest = (biomass_export_mg.get(m, 0.0) if include_uptake else 0.0) * alloc
        leaching = leach[m] * duration_yr * alloc
        erosion = (soil[m] or 0.0) * soil_eroded_kg * a * f_er * alloc  # mg/kg * kg = mg
        balance = m_agro - harvest - leaching - erosion
        out[m] = HeavyMetalResult(
            allocation=alloc,
            leaching_groundwater_kg=leaching * (1 - drained_fraction) * MG_TO_KG,
            leaching_surface_water_kg=leaching * drained_fraction * MG_TO_KG,
            erosion_kg=erosion * MG_TO_KG,
            soil_kg=balance * MG_TO_KG,
        )
    return out


# --------------------------------------------------------------------------- land occupation (2.9.2)
def land_occupation_m2a(duration_months: float = 12.0, area_ha: float = 1.0) -> float:
    """Land occupation [m2*a] for every crop, flooded rice included; 12 months by default."""
    return area_ha * 10_000 * duration_months / 12


# --------------------------------------------------------------------------- seeds: master data
@dataclass(frozen=True)
class IntermediateExchange:
    uuid: str
    name: str
    unit: str


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def load_intermediate_exchanges(path: str | Path, sheet: str = "IntermediateExchanges",
                                name_col: str = "name", id_col: str = "id",
                                unit_col: str = "unitName") -> dict[str, IntermediateExchange]:
    """Load ecoinvent Master Data intermediate exchanges, keyed by exact English name.

    Accepts the ecoSpold2 master-data XML (IntermediateExchanges.xml, namespace-agnostic)
    or a table (.xlsx sheet / .csv) with the given column names. Adjust the column names
    to the actual v3.12 file.
    """
    path = Path(path)
    rows: list[tuple[str, str, str]] = []
    if path.suffix.lower() == ".xml":
        for el in ET.parse(path).getroot().iter():
            if _local(el.tag) != "intermediateExchange":
                continue
            name = unit = ""
            for child in el:
                lang = child.get("{http://www.w3.org/XML/1998/namespace}lang", "en")
                if _local(child.tag) == "name" and lang == "en":
                    name = (child.text or "").strip()
                elif _local(child.tag) == "unitName" and lang == "en":
                    unit = (child.text or "").strip()
            rows.append((el.get("id", ""), name, unit))
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True)[sheet]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        i_id, i_name, i_unit = (header.index(c) for c in (id_col, name_col, unit_col))
        rows = [(str(r[i_id]), str(r[i_name]).strip(), str(r[i_unit])) for r in it if r[i_name]]
    else:
        with open(path, newline="", encoding="utf-8") as f:
            rows = [(r[id_col], r[name_col].strip(), r[unit_col]) for r in csv.DictReader(f)]
    return {name: IntermediateExchange(uuid, name, unit) for uuid, name, unit in rows if name}


def normalise_crop_key(name: str) -> str:
    """Canonical crop key: lower case, separators -> '_'."""
    key = re.sub(r"[\s\-/,.]+", "_", name.strip().lower())
    return re.sub(r"_+", "_", key).strip("_")


@dataclass(frozen=True)
class SeedMappingRow:
    exchange_name: str | None
    status: str          # confirmed | no_exchange | proposed | to_decide


USABLE_STATUS = ("confirmed", "no_exchange")


def load_seed_mapping(path: str | Path = DATA_DIR / "seed_exchange_mapping.csv") -> dict[str, SeedMappingRow]:
    """crop_key -> mapping row, from the reviewed mapping CSV.

    status: "confirmed" (use the exchange), "no_exchange" (crop has no seed input),
    "proposed" / "to_decide" (not yet reviewed: the run stops).
    """
    with open(path, newline="", encoding="utf-8") as f:
        return {normalise_crop_key(r["crop_key"]):
                SeedMappingRow(r["seed_exchange_name"].strip() or None, r["status"].strip())
                for r in csv.DictReader(f)}


def validate_seed_mapping(mapping: Mapping[str, SeedMappingRow],
                          exchanges: Mapping[str, IntermediateExchange]) -> list[str]:
    """Mapped exchange names that do not exist in the master data (should be empty)."""
    return sorted({row.exchange_name for row in mapping.values()
                   if row.exchange_name and row.exchange_name not in exchanges})


def resolve_seed_exchange(crop: str, mapping: Mapping[str, SeedMappingRow],
                          exchanges: Mapping[str, IntermediateExchange],
                          expected_unit: str | None = None) -> IntermediateExchange | None:
    """Seed/seedling exchange for a crop, failing loudly instead of dropping the input.

    Returns None only for crops explicitly marked "no_exchange". Raises if the crop is not
    in the mapping, not yet reviewed, not in the master data, or if ``expected_unit``
    (the unit of the amount in the template, e.g. "kg" or "unit") does not match.
    """
    key = normalise_crop_key(crop)
    if key not in mapping:
        raise KeyError(f"No seed mapping for crop '{crop}' (key '{key}'). "
                       f"Candidates: {suggest_seed_exchanges(crop, exchanges)}")
    row = mapping[key]
    if row.status not in USABLE_STATUS:
        raise ValueError(f"Seed mapping for '{crop}' has status '{row.status}': review it first")
    if row.status == "no_exchange":
        return None
    if row.exchange_name not in exchanges:
        raise KeyError(f"Mapped exchange '{row.exchange_name}' for '{crop}' not in master data")
    ex = exchanges[row.exchange_name]
    if expected_unit is not None and ex.unit != expected_unit:
        raise ValueError(f"Unit mismatch for '{crop}': exchange in {ex.unit}, amount in {expected_unit}")
    return ex


def suggest_seed_exchanges(crop: str, exchanges: Iterable[str], n: int = 5) -> list[str]:
    """Candidate seed/seedling exchange names for building the mapping (review by hand).

    Only for proposing candidates: the run itself never picks a fuzzy match.
    """
    words = normalise_crop_key(crop).split("_")
    pool = [e for e in exchanges if re.search(r"seed|seedling|sowing|planting", e, re.I)]
    hits = [e for e in pool if all(w in e.lower() for w in words)]
    return hits[:n] or difflib.get_close_matches(f"{crop} seed, for sowing", pool, n=n, cutoff=0.4)
