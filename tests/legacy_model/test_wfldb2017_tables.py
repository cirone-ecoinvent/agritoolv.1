"""Decision L22: country soil, precipitation and irrigation water source tables of defaultTables.py
from the WFLDB 2017 background parameters (ecoinvent 3 report, 04 May 2017). The manure shares of
sheet 3 were replaced by FAOSTAT 2021-2023 (decision L29, test_manure_defaults.py)."""
from pathlib import Path

import pytest

import dataloader
import defaultTables as t
from models.irrigationmodel import WaterUseType

WORKBOOK = (Path(__file__).resolve().parents[2] / "Literature"
            / "ecoinvent 3 report_WFLDB_BackgroundParameters_2017.xlsx")
COUNTRY_TABLES = ["ANNUAL_PRECIPITATION_PER_COUNTRY", "CLAY_CONTENT_PER_COUNTRY", "SAND_CONTENT_PER_COUNTRY",
                  "SOIL_CARBON_CONTENT_PER_COUNTRY", "SOIL_WITH_PH_UNDER_OR_7_PER_COUNTRY",
                  "YEARLY_PRECIPITATION_AS_SNOW_PER_COUNTRY", "IRR_WATERUSE_RATIO_PER_COUNTRY"]
NOT_IN_WFLDB = {"MA", "CS"}


@pytest.mark.parametrize("name", COUNTRY_TABLES)
def test_tables_cover_the_39_countries_of_the_tool(name):
    assert set(getattr(t, name)) == set(dataloader.table("luc_blonk_country_expansion.csv"))


def test_water_shares_keep_the_enum_keys_and_sum_to_one():
    for shares in t.IRR_WATERUSE_RATIO_PER_COUNTRY.values():
        assert set(shares) == set(WaterUseType) and sum(shares.values()) == pytest.approx(1.0, abs=1e-9)


def test_spot_values_from_the_workbook():
    # sheet '1) Agr. system parameters': P12 (FR clay), W17 (IT precipitation), D7 (AR sandy soils)
    assert t.CLAY_CONTENT_PER_COUNTRY["FR"] == 0.304420218037661
    assert t.ANNUAL_PRECIPITATION_PER_COUNTRY["IT"] == 930.193617321
    assert t.SAND_CONTENT_PER_COUNTRY["AR"] == 0.04018014391385297


def test_china_humus_is_the_documented_estimate():
    # workbook cell J13 is 0; WFLDB Documentation 2017 p. 14 note 3: estimate of 2 % for China
    assert t.SOIL_CARBON_CONTENT_PER_COUNTRY["CN"] == 0.02


@pytest.fixture(scope="module")
def workbook():
    if not WORKBOOK.is_file():
        pytest.skip("WFLDB 2017 workbook not in Literature/")
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # "Unknown extension is not supported"
        return openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=False)


def _columns(ws, first_col):
    return {ws.cell(3, c).value: c for c in range(first_col, ws.max_column + 1) if ws.cell(3, c).value}


def test_agricultural_parameters_equal_the_workbook(workbook):
    ws = workbook["1) Agr. system parameters"]
    cols = _columns(ws, 4)
    rows = {"ANNUAL_PRECIPITATION_PER_COUNTRY": 17, "CLAY_CONTENT_PER_COUNTRY": 12, "SAND_CONTENT_PER_COUNTRY": 7,
            "SOIL_CARBON_CONTENT_PER_COUNTRY": 13, "SOIL_WITH_PH_UNDER_OR_7_PER_COUNTRY": 14,
            "YEARLY_PRECIPITATION_AS_SNOW_PER_COUNTRY": 18}
    for name, row in rows.items():
        for country, value in getattr(t, name).items():
            if country in NOT_IN_WFLDB or (name == "SOIL_CARBON_CONTENT_PER_COUNTRY" and country == "CN"):
                continue
            assert value == ws.cell(row, cols[country]).value, (name, country)


def test_clay_is_the_soil_type_mix_of_documentation_tab6(workbook):
    ws = workbook["1) Agr. system parameters"]
    clay_of_soil_type = {5: 0.30, 6: 0.50, 7: 0.04, 8: 0.35, 9: 0.35, 10: 0.35}  # HAC, LAC, sandy, spodic, volcanic, wetland
    for country, col in _columns(ws, 4).items():
        if country in t.CLAY_CONTENT_PER_COUNTRY:
            mix = sum(ws.cell(row, col).value * clay for row, clay in clay_of_soil_type.items())
            assert t.CLAY_CONTENT_PER_COUNTRY[country] == pytest.approx(mix, abs=1e-12), country


def test_water_sources_equal_the_workbook(workbook):
    ws = workbook["2) Irrigation water sources"]
    cols = _columns(ws, 2)
    for country, shares in t.IRR_WATERUSE_RATIO_PER_COUNTRY.items():
        if country not in NOT_IN_WFLDB:
            for member, row in ((WaterUseType.groundwater, 4), (WaterUseType.surfacewater, 5),
                                (WaterUseType.nonconventional, 6)):
                assert shares[member] == ws.cell(row, cols[country]).value, (country, member)


def test_erodibility_is_the_soil_type_mix_of_documentation_tab6():
    # decision L27: K = sum of the soil-type shares (sheet 1 rows 5-10) x K-factor ** of Tab. 6 (p. 14)
    k_factors = dataloader.table("wfldb2017_soil_type_k_factors.csv")
    assert {t: r["k_factor_panagos2014_t_h_per_mj_mm"] for t, r in k_factors.items()} == {
        "hac": 0.0347, "lac": 0.01325, "sandy": 0.0122, "spodic": 0.02735, "volcanic": 0.0385, "wetland": 0.0032}
    parameters = dataloader.table("wfldb2017_agr_system_parameters.csv")
    for country, k in t.SOIL_ERODIBILITY_FACTOR_PER_COUNTRY.items():
        shares = {soil: parameters[country][soil + "_soils_ratio"] for soil in k_factors}
        assert sum(shares.values()) == pytest.approx(1.0, abs=1e-9), country
        assert k == pytest.approx(sum(shares[s] * k_factors[s]["k_factor_panagos2014_t_h_per_mj_mm"] for s in shares))
    assert t.SOIL_ERODIBILITY_FACTOR_PER_COUNTRY["IT"] == pytest.approx(0.0346, abs=5e-5)
    assert set(t.SOIL_ERODIBILITY_FACTOR_PER_COUNTRY) == set(t.CLAY_CONTENT_PER_COUNTRY) - NOT_IN_WFLDB


def test_k_factor_from_clay_column_matches_the_tab6_formula():
    for row in dataloader.table("wfldb2017_soil_type_k_factors.csv").values():
        x = 100.0 * row["clay_content_ratio"]
        assert row["k_factor_from_clay_t_h_per_mj_mm"] == pytest.approx(
            -0.00002217 * x * x + 0.00171117 * x + 0.00572691, abs=5e-5)


@pytest.mark.parametrize("country", sorted(NOT_IN_WFLDB))
def test_countries_without_wfldb_soil_types_keep_the_texture_lookup(country):
    from defaultGeneration import DefaultValuesWrapper, DEFAULTS_VALUES_GENERATORS
    defaults = DefaultValuesWrapper({"country": country}, DEFAULTS_VALUES_GENERATORS)
    assert defaults["soil_erodibility_factor"] == t.SOIL_ERODIBILITY_FACTOR_PER_SOIL_TEXTURE[defaults["soil_texture"]]
