"""Decision L43: default total pesticides (g a.i./ha) from the WFLDB 3.1 Documentation 2017, row 19
'Total amount of active ingredient (a.i.), any kind of pesticides' of the sheets 2a-2c
(data/wfldb2017_pesticides_total_ai.csv), mapped to the tool crops by
data/wfldb2017_pesticides_crop_mapping.csv; GLO = mean of the direct datasets of the crop."""
import re
from pathlib import Path

import pytest

import dataloader
from defaultGeneration import DefaultValuesWrapper, DEFAULTS_VALUES_GENERATORS
from defaultMatrixTotalPesticides import TOTAL_PESTICIDES_PER_CROP_PER_COUNTRY as PESTICIDES

WORKBOOK = (Path(__file__).resolve().parents[2] / "Literature"
            / "ecoinvent 3 report_WFLDB_Documentation_2017.xlsx")
SHEETS = ("2a) Main inputs_Arable", "2b) Main inputs_Horticulture", "2c) Main inputs_Perennials")
LEGACY_TAG = "legacy, undocumented, provisional"


def _default(**inputs):
    return DefaultValuesWrapper(inputs, DEFAULTS_VALUES_GENERATORS)["pest_total"]


def _datasets():
    return {row["wfldb_dataset_code"]: row for row in dataloader.rows("wfldb2017_pesticides_total_ai.csv")}


def _mapping():
    return dataloader.rows("wfldb2017_pesticides_crop_mapping.csv")


def test_transcription_has_the_104_datasets_of_the_three_sheets():
    datasets = _datasets()
    assert len(datasets) == 104
    assert {row["sheet"] for row in datasets.values()} == set(SHEETS)
    assert {row["data_level"] for row in datasets.values()} == {"L0", "L1", "L3", "L4"}


def test_country_is_the_suffix_of_the_dataset_code():
    for code, row in _datasets().items():
        assert re.match(r"^[A-Za-z]+?%s(_.*|pross)?$" % row["country"], code)
        name_country = re.search(r"/([A-Z]{2}) U$", row["wfldb_dataset"]).group(1)
        # the name of the Malaysian palm dataset ends with /ID in the workbook (row 3)
        assert name_country == row["country"] or (code, name_country) == ("PalmFruitBunchMY", "ID")


def test_spot_values_of_the_transcription():
    datasets = _datasets()
    almond = datasets["AlmondCN"]
    assert (almond["total_ai_kg_per_ha"], almond["data_level"], almond["cell"]) == ("11.1", "L1", "C19")
    assert datasets["PearAR"]["total_ai_kg_per_ha"] == "13.78"
    assert datasets["StrawberryES"]["total_ai_kg_per_ha"] == "26.3"
    assert datasets["CoffeeIN_Rob"]["total_ai_kg_per_ha"] == "0"


@pytest.fixture(scope="module")
def workbook():
    if not WORKBOOK.is_file():
        pytest.skip("WFLDB 2017 Documentation workbook not in Literature/")
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # "Unknown extension is not supported"
        return openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=False)


def test_transcription_equals_the_workbook(workbook):
    from openpyxl.utils import get_column_letter
    expected = {}
    for name in SHEETS:
        ws = workbook[name]
        assert ws.cell(19, 1).value.startswith("Total amount of active ingredient (a.i.)")
        assert ws.cell(19, 2).value == "kg/ha"
        for col in range(3, ws.max_column + 1, 2):
            code = ws.cell(4, col).value
            if code:
                expected[code] = (ws.cell(3, col).value, float(ws.cell(19, col).value), ws.cell(19, col + 1).value,
                                  name, "%s19" % get_column_letter(col))
    transcribed = {code: (row["wfldb_dataset"], float(row["total_ai_kg_per_ha"]), row["data_level"], row["sheet"],
                          row["cell"]) for code, row in _datasets().items()}
    assert transcribed == expected


def test_mapping_covers_every_dataset_with_tool_crops():
    crops = set(dataloader.table("crop_types.csv", numeric=False))
    mapping = _mapping()
    assert {row["wfldb_dataset_code"] for row in mapping} == set(_datasets())
    assert {row["use"] for row in mapping} == {"direct", "other_grade", "not_used"}
    assert all(row["crop"] in crops for row in mapping if row["use"] != "not_used")
    assert {row["wfldb_dataset_code"] for row in mapping if row["use"] == "not_used"} == {"StrawberryCH_GH", "VanillaMG"}


def test_mapping_is_signed():
    assert all(row["status"] == "confirmed" and row["reviewed_by"].strip() for row in _mapping())


def test_values_and_glo_by_hand():
    # kg a.i./ha x 1000; legacy values in the comments
    assert PESTICIDES["almond"] == pytest.approx({"CN": 11100.0, "US": 10920.0, "GLO": (11100.0 + 10920.0) / 2})
    assert PESTICIDES["pear"]["AR"] == pytest.approx(13780.0)  # 866.4
    assert PESTICIDES["strawberry_fresh"]["ES"] == pytest.approx(26300.0)  # 217792.35
    assert PESTICIDES["coffee_robusta"]["VN"] == pytest.approx(500.0)  # 9631.6
    assert PESTICIDES["rice"]["IN"] == pytest.approx(1500.0)  # 1498.06, rice_IN snapshot


def test_grades_and_species():
    fresh, processing = PESTICIDES["tomato_fresh"], PESTICIDES["tomato_processing"]
    assert (fresh["ES"], processing["ES"]) == pytest.approx((32300.0, 19520.0))
    assert (fresh["IT"], processing["IT"]) == pytest.approx((21800.0, 21800.0))  # processing grade only
    assert (fresh["NL"], processing["NL"]) == pytest.approx((45690.0, 45690.0))  # fresh grade only
    assert fresh["GLO"] == pytest.approx((45690.0 + 32300.0 + 39730.0) / 3)
    assert processing["GLO"] == pytest.approx((19520.0 + 21800.0) / 2)
    fresh, processing = PESTICIDES["orange_fresh"], PESTICIDES["orange_processing"]
    assert (fresh["US"], processing["US"]) == pytest.approx((11000.0, 14450.0))
    assert (fresh["BR"], processing["BR"]) == pytest.approx((6900.0, 6900.0))  # processing grade only
    assert processing["GLO"] == pytest.approx((6900.0 + 14450.0) / 2)
    arabica, robusta = PESTICIDES["coffee_arabica"], PESTICIDES["coffee_robusta"]
    assert (arabica["BR"], robusta["BR"]) == pytest.approx((7820.0, 5760.0))
    assert (arabica["IN"], robusta["IN"]) == pytest.approx((3850.0, 0.0))
    assert (arabica["ID"], robusta["ID"]) == pytest.approx((230.0, 230.0))  # Robusta only
    assert (arabica["CO"], robusta["CO"]) == pytest.approx((4310.0, 4310.0))  # Arabica only
    assert arabica["GLO"] == pytest.approx((7820.0 + 1970.0 + 4310.0 + 3850.0) / 4)  # with HN
    assert robusta["GLO"] == pytest.approx((5760.0 + 0.0 + 230.0 + 500.0) / 4)


def test_strawberry_ch_takes_the_heated_greenhouse():
    assert PESTICIDES["strawberry_fresh"]["CH"] == pytest.approx(45770.0)  # non-heated 27460 not used


def test_new_countries():
    assert PESTICIDES["palmtree"] == pytest.approx({"MY": 1700.0, "ID": 2600.0, "GLO": 2150.0})
    assert PESTICIDES["maizegrain"] == pytest.approx({"AR": 130.0, "GLO": 130.0})
    assert PESTICIDES["sugarbeet"]["DE"] == pytest.approx(6960.0)
    assert PESTICIDES["onion"]["IN"] == pytest.approx(3830.0)


def test_countries_without_wfldb_dataset():
    # strawberry MX and MA repeated the legacy GLO: removed, they take the new GLO
    assert _default(crop="strawberry_fresh", country="MX") == pytest.approx((26300.0 + 45770.0 + 41800.0) / 3)
    assert PESTICIDES["tomato_fresh"]["CH"] == 27206.5388274181  # legacy kept


def test_legacy_table_keeps_only_the_rows_without_wfldb_dataset():
    rows = dataloader.rows("default_pesticides.csv")
    assert len(rows) == 36 and all(row["note"].endswith(LEGACY_TAG) for row in rows)
    mapped = {row["crop"] for row in _mapping() if row["use"] != "not_used"}
    assert {(row["crop"], row["country"]) for row in rows if row["crop"] in mapped} == {
        ("tomato_fresh", "CH"), ("tomato_processing", "CH")}
